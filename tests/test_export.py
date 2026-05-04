"""Tests for LaTeX and DOCX exporters."""
from __future__ import annotations

import textwrap
from pathlib import Path

import pytest

from paper_agent.agents.schema import Suggestion, _uid
from paper_agent.export import DocxExporter, LaTeXExporter
from paper_agent.parsers import LaTeXParser, WordParser


SAMPLE_TEX = textwrap.dedent(r"""
    \documentclass{article}
    \usepackage{hyperref}
    \title{Foo}
    \begin{document}
    \begin{abstract}
    We achieve 92.4\% on Libero.
    \end{abstract}
    \section{Introduction}
    In recent years, robotics has progressed rapidly. We address this gap.
    \end{document}
""").strip()


def _mk_sug(pid: str, before: str, after: str, status: str = "accepted",
            rationale: str = "remove ai-flavor opener", flag: bool = False) -> Suggestion:
    return Suggestion(
        id=_uid("sug"), todo_id="t1", paragraph_id=pid, section_id="",
        category="language", severity="info",
        before=before, after=after, rationale=rationale, flag_for_user=flag,
        status=status,
    )


def test_latex_export_writes_accepted_as_plain_text(tmp_path: Path):
    tree = LaTeXParser().parse(SAMPLE_TEX)
    intro = tree.sections[0]
    p = intro.paragraphs[0]
    assert "In recent years, robotics" in p.raw_source

    sug_accept = _mk_sug(p.id, "In recent years, robotics", "Robotics", "accepted")
    sug_pending = _mk_sug(p.id, "We address this gap.", "We close this gap.", "pending")
    sug_reject = _mk_sug(p.id, "rapidly", "fast", "rejected")

    out = tmp_path / "out.tex"
    summary = LaTeXExporter().export(tree, [sug_accept, sug_pending, sug_reject], str(out))

    written = out.read_text(encoding="utf-8")
    # accepted -> applied as plain text replacement
    assert "Robotics has progressed rapidly" in written
    assert "In recent years, robotics" not in written
    # NO \replaced wrapper, NO changes package
    assert "\\replaced" not in written
    assert "{changes}" not in written
    # pending NOT in output (rejected ditto)
    assert "We close this gap." not in written
    assert "We address this gap." in written  # original preserved
    # tally
    assert summary.accepted == 1 and summary.pending == 1 and summary.rejected == 1


def test_latex_export_flags_pending_for_user(tmp_path: Path):
    """Pending suggestions should be tallied so the UI can warn the user
    that their pending decisions are not being exported."""
    tree = LaTeXParser().parse(SAMPLE_TEX)
    p = tree.sections[0].paragraphs[0]
    sug_pending_flagged = _mk_sug(p.id, "rapidly", "quickly", "pending", flag=True)
    sug_pending_plain = _mk_sug(p.id, "address this gap", "close this gap", "pending")
    summary = LaTeXExporter().export(
        tree, [sug_pending_flagged, sug_pending_plain], str(tmp_path / "x.tex")
    )
    assert summary.pending == 2
    assert sug_pending_flagged.id in summary.flagged_pending
    assert sug_pending_plain.id not in summary.flagged_pending


def test_latex_skipped_when_before_not_found(tmp_path: Path):
    tree = LaTeXParser().parse(SAMPLE_TEX)
    p = tree.sections[0].paragraphs[0]
    bad = _mk_sug(p.id, "this exact phrase is not present", "X", "accepted")
    summary = LaTeXExporter().export(tree, [bad], str(tmp_path / "x.tex"))
    assert summary.skipped_not_found == 1


def test_docx_export_pending_writes_track_changes(tmp_path: Path):
    pytest.importorskip("docx")
    pytest.importorskip("lxml")
    from docx import Document

    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("In recent years, robotics has progressed rapidly. We address this gap.")
    doc.save(src)

    tree = WordParser().parse_file(str(src))
    p = next(p for p in tree.all_paragraphs() if "recent years" in p.text)

    accepted = _mk_sug(p.id, "In recent years, robotics", "Robotics", "accepted")
    pending = _mk_sug(p.id, "We address this gap.", "We close this gap.", "pending")

    out = tmp_path / "out.docx"
    DocxExporter().export(tree, [accepted, pending], str(src), str(out))

    # Check the resulting docx contains <w:ins> and <w:del> for the pending edit
    import zipfile

    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")

        # accepted text appears plainly (may be split across runs)
        assert "Robotics" in xml and "has progressed rapidly" in xml
        # original "In recent years, " is gone (accepted edit applied directly)
        assert "In recent years" not in xml
        # pending appears as track changes
        assert "<w:ins" in xml and "<w:del" in xml
        # the new text lives inside <w:ins>...</w:ins>
        ins_start = xml.find("<w:ins")
        ins_end = xml.find("</w:ins>")
        assert "We close this gap." in xml[ins_start:ins_end]
        # the original text lives inside <w:del>...</w:del>
        del_start = xml.find("<w:del ")
        del_end = xml.find("</w:del>")
        assert "We address this gap." in xml[del_start:del_end]


def test_latex_export_does_not_leak_ai_rationale(tmp_path: Path):
    """The exported .tex must be clean of AI metadata.

    Even when every suggestion carries a verbose rationale, NONE of it
    should end up in the file. The Web UI is the only place that surfaces
    AI opinions; the exported manuscript looks human-edited.
    """
    tree = LaTeXParser().parse(SAMPLE_TEX)
    p = tree.sections[0].paragraphs[0]
    sug = _mk_sug(
        p.id, "In recent years, robotics", "Robotics",
        status="accepted",
        rationale="Drop empty 'in recent years' opener; ai-flavor.",
    )
    out = tmp_path / "out.tex"
    LaTeXExporter().export(tree, [sug], str(out))

    written = out.read_text(encoding="utf-8")
    # The replacement is applied as plain text, no markers.
    assert "Robotics has progressed" in written
    # No tag, no rationale text, no `% [AI...]` style comments.
    assert "[AI/" not in written
    assert "[AI:" not in written
    assert "ai-flavor" not in written
    assert "Drop empty" not in written


def test_docx_export_accepted_is_clean_no_comments(tmp_path: Path):
    """Accepted/modified edits in the .docx must NOT carry any AI metadata.

    No `word/comments.xml` part, no commentRange markers, no rationale text
    leaks from the suggestions.
    """
    pytest.importorskip("docx")
    from docx import Document

    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "In recent years, robotics has progressed rapidly. We address this gap."
    )
    doc.save(src)

    tree = WordParser().parse_file(str(src))
    p = next(p for p in tree.all_paragraphs() if "recent years" in p.text)
    accepted = _mk_sug(
        p.id, "In recent years, robotics", "Robotics",
        status="accepted", rationale="Drop ai-flavor opener.",
    )
    modified = _mk_sug(
        p.id, "rapidly", "quickly",
        status="modified", rationale="Tighten verb.",
    )

    out = tmp_path / "out.docx"
    DocxExporter().export(tree, [accepted, modified], str(src), str(out))

    import zipfile

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "word/comments.xml" not in names, names
        doc_xml = z.read("word/document.xml").decode("utf-8")

    assert "<w:commentRangeStart" not in doc_xml
    assert "<w:commentRangeEnd" not in doc_xml
    assert "<w:commentReference" not in doc_xml
    assert "Drop ai-flavor opener" not in doc_xml
    assert "Tighten verb" not in doc_xml
    # Edits are still applied as plain text.
    assert "Robotics" in doc_xml
    assert "quickly" in doc_xml
    # And accepted/modified do NOT produce track-change marks.
    assert "<w:ins" not in doc_xml
    assert "<w:del" not in doc_xml


def test_docx_export_pending_carries_ai_comment(tmp_path: Path):
    """Pending edits in the .docx should appear as track-changes AND carry
    the AI rationale as a Word comment anchored on the inserted text.

    Word's Comments pane will then show "[category/severity] rationale"
    next to each unresolved <w:ins>...</w:ins> region.
    """
    pytest.importorskip("docx")
    from docx import Document

    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "In recent years, robotics has progressed rapidly. We address this gap."
    )
    doc.save(src)

    tree = WordParser().parse_file(str(src))
    p = next(p for p in tree.all_paragraphs() if "recent years" in p.text)
    accepted = _mk_sug(
        p.id, "In recent years, robotics", "Robotics",
        status="accepted", rationale="Drop ai-flavor opener.",
    )
    pending = _mk_sug(
        p.id, "We address this gap.", "We close this gap.",
        status="pending", rationale="`address` is overused; prefer `close`.",
    )

    out = tmp_path / "out.docx"
    DocxExporter().export(tree, [accepted, pending], str(src), str(out))

    import zipfile

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "word/comments.xml" in names, names
        doc_xml = z.read("word/document.xml").decode("utf-8")
        comments_xml = z.read("word/comments.xml").decode("utf-8")

    # The pending edit is rendered as track changes and the AI rationale
    # rides as a real Word comment.
    assert "<w:ins" in doc_xml and "<w:del" in doc_xml
    assert "<w:commentRangeStart" in doc_xml
    assert "<w:commentRangeEnd" in doc_xml
    assert "<w:commentReference" in doc_xml

    # Comment body carries category/severity-tagged rationale.
    assert "[language/info]" in comments_xml
    assert "address` is overused" in comments_xml
    # The accepted edit's rationale must NOT leak into comments.
    assert "Drop ai-flavor opener" not in comments_xml

    # Exactly one comment range pair was emitted (one pending edit).
    assert doc_xml.count("<w:commentRangeStart") == 1
    assert doc_xml.count("<w:commentRangeEnd") == 1


def test_docx_export_pending_without_rationale_skips_comment(tmp_path: Path):
    """If a pending suggestion has no rationale we still write the
    track-change but skip the Word comment so the doc isn't polluted with
    empty comment markers."""
    pytest.importorskip("docx")
    from docx import Document

    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_paragraph("Some text that needs editing.")
    doc.save(src)

    tree = WordParser().parse_file(str(src))
    p = next(p for p in tree.all_paragraphs() if "Some text" in p.text)
    pending = _mk_sug(
        p.id, "Some text", "Different text",
        status="pending", rationale="",
    )

    out = tmp_path / "out.docx"
    DocxExporter().export(tree, [pending], str(src), str(out))

    import zipfile

    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        doc_xml = z.read("word/document.xml").decode("utf-8")

    # Track-change present, comments part absent.
    assert "<w:ins" in doc_xml
    assert "word/comments.xml" not in names
    assert "<w:commentRangeStart" not in doc_xml


def test_docx_export_uses_custom_author(tmp_path: Path):
    pytest.importorskip("docx")
    from docx import Document

    src = tmp_path / "in.docx"
    doc = Document()
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Some text that needs editing.")
    doc.save(src)

    tree = WordParser().parse_file(str(src))
    p = next(p for p in tree.all_paragraphs() if "Some text" in p.text)
    pending = _mk_sug(p.id, "Some text", "Different text", "pending")

    out = tmp_path / "out.docx"
    DocxExporter().export(
        tree, [pending], str(src), str(out), author="Alice Reviewer"
    )

    import zipfile

    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert 'w:author="Alice Reviewer"' in xml
    assert "Paper-Agent AI" not in xml
